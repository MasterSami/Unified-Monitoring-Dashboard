"""JSON API for the What-If capacity view (``/api/v1/whatif``).

The page itself is served by :mod:`app.routers.pages` as the third Capacity
tab; this module only exposes the data the page talks to.
"""

from __future__ import annotations

import io
from copy import copy
from datetime import datetime, timezone

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import Response
from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.db import get_db
from app.export_xlsx import build_workbook
from app.models import WhatIfScenario
from app.whatif.engine import op_label, simulate, target_catalog
from app.whatif.schemas import ScenarioDefinition, SimulationRequest

_API = APIRouter(prefix="/api/v1/whatif", tags=["whatif"])

_DISCLAIMER = "Simulation - a straight-line projection of the current trend, not a prediction."


@_API.get("/targets")
def targets(
    q: str = "",
    kind: str = "all",
    host_id: int | None = Query(default=None, ge=1),
    db: Session = Depends(get_db),
):
    return {"targets": target_catalog(db, q, kind, host_id)}


@_API.post("/simulate")
def simulate_api(payload: SimulationRequest, db: Session = Depends(get_db)):
    try:
        return simulate(db, payload)
    except ValueError as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc


def _scenario_out(row: WhatIfScenario) -> dict:
    return {
        "id": row.id, "name": row.name, "description": row.description,
        "target": row.target_json.get("id", ""),
        "targets": row.target_json.get("targets", []),
        "ops": row.ops_json, "updated_at": row.updated_at,
    }


@_API.get("/scenarios")
def list_scenarios(db: Session = Depends(get_db)):
    rows = db.scalars(select(WhatIfScenario).order_by(WhatIfScenario.updated_at.desc())).all()
    return [_scenario_out(r) for r in rows]


@_API.post("/scenarios")
def save_scenario(payload: ScenarioDefinition, db: Session = Depends(get_db)):
    row = WhatIfScenario(
        name=payload.name, description=payload.description,
        target_json={"id": payload.target, "targets": payload.targets},
        ops_json=[op.model_dump(mode="json", exclude_none=True) for op in payload.ops],
        schema_version=1,
    )
    db.add(row)
    db.commit()
    db.refresh(row)
    return _scenario_out(row)


@_API.put("/scenarios/{scenario_id}")
def update_scenario(scenario_id: int, payload: ScenarioDefinition, db: Session = Depends(get_db)):
    row = db.get(WhatIfScenario, scenario_id)
    if row is None:
        raise HTTPException(404, "scenario not found")
    row.name, row.description = payload.name, payload.description
    row.target_json = {"id": payload.target, "targets": payload.targets}
    row.ops_json = [op.model_dump(mode="json", exclude_none=True) for op in payload.ops]
    db.commit()
    db.refresh(row)
    return _scenario_out(row)


@_API.delete("/scenarios/{scenario_id}")
def delete_scenario(scenario_id: int, db: Session = Depends(get_db)):
    row = db.get(WhatIfScenario, scenario_id)
    if row is None:
        raise HTTPException(404, "scenario not found")
    db.delete(row)
    db.commit()
    return {"deleted": scenario_id}


def _describe_op(op: dict) -> str:
    kind = op.get("kind", "")
    label = op_label(kind)
    if kind in ("growth_multiplier", "utilization_multiplier"):
        return f"{label}: x{op.get('factor', 1)}"
    if kind == "threshold":
        return f"{label}: {op.get('threshold', 90)}%"
    if kind == "horizon_date":
        return f"{label}: {op.get('horizon_date', '')}"
    if kind == "move_workload":
        return f"{label}: {op.get('size_gb', op.get('value', 0))} GB, {op.get('daily_gb', 0)} GB/day to {op.get('destination', '')}"
    if kind == "resize" and op.get("absolute_total") is not None:
        return f"{label}: to {op['absolute_total']}"
    unit = {"add_daily": "/day", "remove_daily": "/day"}.get(kind, "")
    return f"{label}: {op.get('value', 0)}{unit}"


@_API.get("/export")
def export_scenario(scenario_id: int, db: Session = Depends(get_db)):
    row = db.get(WhatIfScenario, scenario_id)
    if row is None:
        raise HTTPException(404, "scenario not found")
    payload = SimulationRequest(
        target=row.target_json.get("id", ""),
        targets=row.target_json.get("targets", []),
        ops=row.ops_json,
    )
    try:
        result = simulate(db, payload)
    except ValueError as exc:
        raise HTTPException(422, str(exc)) from exc

    columns = ["Target", "Metric", "Baseline", "Scenario", "Delta"]
    rows: list[list] = []
    for pt in result["per_target"]:
        b, s = pt["baseline"], pt["scenario"]
        name = f"{pt['hostname']} {pt['kind']} {pt['subject']}".strip()
        unit = b.get("unit", "GB")
        rows += [
            [name, "Current %", b["current_pct"], s["current_pct"], round(s["current_pct"] - b["current_pct"], 2)],
            [name, f"Total ({unit})", b["total_value"], s["total_value"],
             None if b["total_value"] is None or s["total_value"] is None else round(s["total_value"] - b["total_value"], 2)],
            [name, "Trend %/day", b["slope_pct_per_day"], s["slope_pct_per_day"], round(s["slope_pct_per_day"] - b["slope_pct_per_day"], 4)],
            [name, f"Days to {s['threshold']:g}%", b["days_to_threshold"], s["days_to_threshold"], pt["days_bought"]],
            [name, "Days to full", b["days_to_full"], s["days_to_full"],
             None if b["days_to_full"] is None or s["days_to_full"] is None else round(s["days_to_full"] - b["days_to_full"], 1)],
            [name, "Classification", b["classification"], s["classification"], ""],
        ]
    data = build_workbook(
        sheet_title="What-If",
        period="Current baseline",
        filters_summary=f"scenario: {row.name}",
        columns=columns,
        rows=rows,
        credit=_DISCLAIMER,
    )

    # Put the scenario's steps above the table so the sheet explains itself.
    wb = load_workbook(io.BytesIO(data))
    ws = wb.active
    header_row = next(r for r in range(1, 12) if ws.cell(r, 1).value == "Target")
    steps = [("Step %d" % i, _describe_op(op)) for i, op in enumerate(row.ops_json, 1)]
    if row.description:
        steps.insert(0, ("Description", row.description))
    ws.insert_rows(header_row, len(steps) + 2)
    ws.cell(header_row, 1, "SCENARIO").font = copy(ws["A1"].font)
    for idx, (label, value) in enumerate(steps, header_row + 1):
        ws.cell(idx, 1, label)
        ws.cell(idx, 2, value)
    out = io.BytesIO()
    wb.save(out)
    stamp = datetime.now(timezone.utc).strftime("%Y%m%d_%H%M%S")
    return Response(
        content=out.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="SAMIX_whatif_{stamp}.xlsx"'},
    )


api_router = _API
